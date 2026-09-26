from __future__ import annotations

import io
import json
import math
import sqlite3
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .database import connect_database

DEVICE_ID = "PT-01"
THRESHOLD_BAR = 0.5
SETTLE_MS = 200
PRETRIGGER_SECONDS = 2.0
TELEMETRY_FRESH_SECONDS = 0.75
MICROS_WRAP = 2**32


class PressureCaptureError(RuntimeError):
    pass


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="milliseconds")


def _parse_stamp(value: str) -> datetime:
    return datetime.fromisoformat(value.replace("Z", "+00:00"))


def ensure_schema(db: sqlite3.Connection) -> None:
    db.executescript(
        """
        CREATE TABLE IF NOT EXISTS pressure_captures(
          id INTEGER PRIMARY KEY AUTOINCREMENT,
          operation_id TEXT NOT NULL,
          device_id TEXT NOT NULL,
          state TEXT NOT NULL,
          started_at TEXT NOT NULL,
          trigger_at TEXT NOT NULL,
          stopped_at TEXT,
          stop_reason TEXT,
          threshold_bar REAL NOT NULL DEFAULT 0.5,
          settle_ms INTEGER NOT NULL DEFAULT 200,
          sample_rate_hz REAL NOT NULL DEFAULT 200,
          seen_above INTEGER NOT NULL DEFAULT 0,
          below_since_rel REAL,
          peak_bar REAL NOT NULL DEFAULT 0,
          peak_time_s REAL,
          sample_count INTEGER NOT NULL DEFAULT 0,
          missing_samples INTEGER NOT NULL DEFAULT 0,
          last_boot_id TEXT,
          last_sequence INTEGER,
          last_extended_us INTEGER,
          boot_wraps INTEGER NOT NULL DEFAULT 0,
          boot_anchor_extended_us INTEGER,
          boot_anchor_rel_s REAL,
          last_sample_rel_s REAL,
          last_sample_received_at TEXT,
          created_by TEXT
        );
        CREATE INDEX IF NOT EXISTS idx_pressure_captures_active
          ON pressure_captures(operation_id,state);
        CREATE TABLE IF NOT EXISTS pressure_capture_samples(
          id INTEGER PRIMARY KEY AUTOINCREMENT,
          capture_id INTEGER NOT NULL,
          sample_index INTEGER NOT NULL,
          boot_id TEXT NOT NULL,
          sequence INTEGER NOT NULL,
          batch_sample_index INTEGER NOT NULL,
          esp_time_us INTEGER NOT NULL,
          relative_time_s REAL NOT NULL,
          pressure_bar REAL NOT NULL,
          voltage_v REAL,
          phase TEXT NOT NULL,
          received_at TEXT NOT NULL,
          UNIQUE(capture_id,boot_id,sequence,batch_sample_index),
          FOREIGN KEY(capture_id) REFERENCES pressure_captures(id)
        );
        CREATE INDEX IF NOT EXISTS idx_pressure_capture_samples_order
          ON pressure_capture_samples(capture_id,sample_index);
        CREATE TABLE IF NOT EXISTS pressure_capture_events(
          id INTEGER PRIMARY KEY AUTOINCREMENT,
          capture_id INTEGER NOT NULL,
          occurred_at TEXT NOT NULL,
          event_type TEXT NOT NULL,
          detail TEXT NOT NULL,
          missing_samples INTEGER NOT NULL DEFAULT 0,
          FOREIGN KEY(capture_id) REFERENCES pressure_captures(id)
        );
        """
    )
    db.commit()


def _event(
    db: sqlite3.Connection,
    capture_id: int,
    event_type: str,
    detail: str,
    missing_samples: int = 0,
) -> None:
    db.execute(
        """INSERT INTO pressure_capture_events(
             capture_id,occurred_at,event_type,detail,missing_samples)
           VALUES(?,?,?,?,?)""",
        (capture_id, utc_now(), event_type, detail, int(max(0, missing_samples))),
    )


def _fresh_telemetry(db: sqlite3.Connection, device_id: str) -> tuple[bool, str]:
    session = db.execute(
        """SELECT status,last_seen FROM edge_sessions
           WHERE device_id=? ORDER BY last_seen DESC LIMIT 1""",
        (device_id,),
    ).fetchone()
    if not session or session["status"] not in {"CONNECTED", "STREAMING"}:
        return False, f"{device_id} Ethernet telemetry is not connected"

    batch = db.execute(
        """SELECT received_at FROM edge_batches
           WHERE device_id=? ORDER BY id DESC LIMIT 1""",
        (device_id,),
    ).fetchone()
    if not batch:
        return False, f"{device_id} has not delivered pressure samples"

    try:
        age = (datetime.now(timezone.utc) - _parse_stamp(batch["received_at"])).total_seconds()
    except (TypeError, ValueError):
        return False, f"{device_id} telemetry timestamp is invalid"
    if age > TELEMETRY_FRESH_SECONDS:
        return False, (
            f"{device_id} pressure telemetry is stale ({age:.2f}s); "
            "ignition command blocked so no test can start without logging"
        )
    return True, ""


def _channel_values(batch: sqlite3.Row | dict) -> tuple[list[float], list[float | None]]:
    channels = json.loads(batch["channels_json"]) if "channels_json" in batch.keys() else batch["channels"]
    pressure_raw = channels.get("pressure_mbar") or []
    voltage_raw = channels.get("voltage_uv") or []
    pressure = [float(value) / 1000.0 for value in pressure_raw]
    voltage: list[float | None] = []
    for index in range(len(pressure)):
        value = voltage_raw[index] if index < len(voltage_raw) else None
        voltage.append(float(value) / 1_000_000.0 if value is not None else None)
    return pressure, voltage


def _copy_pretrigger(
    db: sqlite3.Connection,
    capture_id: int,
    device_id: str,
    trigger_at: datetime,
) -> None:
    cutoff = (trigger_at - timedelta(seconds=PRETRIGGER_SECONDS + 0.5)).isoformat(
        timespec="milliseconds"
    )
    rows = db.execute(
        """SELECT * FROM edge_batches
           WHERE device_id=? AND received_at>=?
           ORDER BY id""",
        (device_id, cutoff),
    ).fetchall()

    selected: list[tuple] = []
    for row in rows:
        pressure, voltage = _channel_values(row)
        count = min(int(row["sample_count"]), len(pressure))
        if not count:
            continue
        period_us = int(row["sample_period_us"])
        received = _parse_stamp(row["received_at"])
        first_wall = received - timedelta(microseconds=period_us * max(0, count - 1))
        for index in range(count):
            wall = first_wall + timedelta(microseconds=period_us * index)
            rel = (wall - trigger_at).total_seconds()
            if -PRETRIGGER_SECONDS <= rel < 0:
                esp_us = (int(row["first_sample_us"]) + period_us * index) % MICROS_WRAP
                selected.append(
                    (
                        str(row["boot_id"]),
                        int(row["sequence"]),
                        index,
                        esp_us,
                        rel,
                        pressure[index],
                        voltage[index],
                        row["received_at"],
                    )
                )

    selected.sort(key=lambda item: item[4])
    for sample_index, item in enumerate(selected):
        db.execute(
            """INSERT OR IGNORE INTO pressure_capture_samples(
                 capture_id,sample_index,boot_id,sequence,batch_sample_index,
                 esp_time_us,relative_time_s,pressure_bar,voltage_v,phase,received_at)
               VALUES(?,?,?,?,?,?,?,?,?,'PRE_TRIGGER',?)""",
            (capture_id, sample_index, *item),
        )

    if not selected:
        return

    last = selected[-1]
    boot_id, sequence, _batch_index, esp_us, rel, pressure_bar, _voltage, received_at = last
    db.execute(
        """UPDATE pressure_captures SET
             sample_count=(SELECT count(*) FROM pressure_capture_samples WHERE capture_id=?),
             peak_bar=?,
             peak_time_s=?,
             last_boot_id=?,
             last_sequence=?,
             last_extended_us=?,
             boot_wraps=0,
             boot_anchor_extended_us=?,
             boot_anchor_rel_s=?,
             last_sample_rel_s=?,
             last_sample_received_at=?
           WHERE id=?""",
        (
            capture_id,
            max(float(x[5]) for x in selected),
            max(selected, key=lambda x: float(x[5]))[4],
            boot_id,
            sequence,
            int(esp_us),
            int(esp_us),
            float(rel),
            float(rel),
            received_at,
            capture_id,
        ),
    )


def start_capture(
    db_path: Path,
    operation_id: str,
    device_id: str = DEVICE_ID,
    actor: str = "TEST_DIRECTOR",
) -> dict:
    db = connect_database(db_path)
    try:
        ensure_schema(db)
        active = db.execute(
            """SELECT * FROM pressure_captures
               WHERE operation_id=? AND state='ACTIVE'
               ORDER BY id DESC LIMIT 1""",
            (operation_id,),
        ).fetchone()
        if active:
            raise PressureCaptureError(
                f"pressure capture {active['id']} is already active; stop it before another ignition"
            )

        fresh, reason = _fresh_telemetry(db, device_id)
        if not fresh:
            raise PressureCaptureError(reason)

        trigger_stamp = utc_now()
        trigger_dt = _parse_stamp(trigger_stamp)
        cursor = db.execute(
            """INSERT INTO pressure_captures(
                 operation_id,device_id,state,started_at,trigger_at,
                 threshold_bar,settle_ms,sample_rate_hz,created_by)
               VALUES(?,?,'ACTIVE',?,?,?,?,200,?)""",
            (
                operation_id,
                device_id,
                trigger_stamp,
                trigger_stamp,
                THRESHOLD_BAR,
                SETTLE_MS,
                actor,
            ),
        )
        capture_id = int(cursor.lastrowid)
        _copy_pretrigger(db, capture_id, device_id, trigger_dt)
        _event(
            db,
            capture_id,
            "CAPTURE_STARTED",
            f"Capture armed before relay command; includes {PRETRIGGER_SECONDS:.1f}s pre-trigger buffer",
        )
        db.commit()
        return capture_status(db, capture_id)
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


def mark_relay_sent(db_path: Path, capture_id: int, relay_at: str | None = None) -> dict:
    """Set T+0 to the actual server-side relay send instant.

    start_capture intentionally arms the historian first. Once the Ethernet
    relay command succeeds, this rebases any samples collected during those few
    milliseconds so exported time zero corresponds to the relay command.
    """
    db = connect_database(db_path)
    try:
        ensure_schema(db)
        row = db.execute(
            "SELECT * FROM pressure_captures WHERE id=?", (capture_id,)
        ).fetchone()
        if not row:
            raise PressureCaptureError("pressure capture not found")
        new_stamp = relay_at or utc_now()
        old_dt = _parse_stamp(row["trigger_at"])
        new_dt = _parse_stamp(new_stamp)
        shift = (old_dt - new_dt).total_seconds()
        db.execute(
            """UPDATE pressure_capture_samples
               SET relative_time_s=relative_time_s+?
               WHERE capture_id=?""",
            (shift, capture_id),
        )
        db.execute(
            """UPDATE pressure_captures SET
                 trigger_at=?,
                 boot_anchor_rel_s=CASE WHEN boot_anchor_rel_s IS NULL THEN NULL ELSE boot_anchor_rel_s+? END,
                 last_sample_rel_s=CASE WHEN last_sample_rel_s IS NULL THEN NULL ELSE last_sample_rel_s+? END,
                 peak_time_s=CASE WHEN peak_time_s IS NULL THEN NULL ELSE peak_time_s+? END,
                 below_since_rel=CASE WHEN below_since_rel IS NULL THEN NULL ELSE below_since_rel+? END
               WHERE id=?""",
            (new_stamp, shift, shift, shift, shift, capture_id),
        )
        _event(db, capture_id, "RELAY_SENT", "Pressure timeline rebased to successful BENCH_LED_SET ON command")
        db.commit()
        return capture_status(db, capture_id)
    finally:
        db.close()


def cancel_capture(db_path: Path, capture_id: int, reason: str) -> None:
    db = connect_database(db_path)
    try:
        ensure_schema(db)
        row = db.execute(
            "SELECT state FROM pressure_captures WHERE id=?", (capture_id,)
        ).fetchone()
        if row and row["state"] == "ACTIVE":
            db.execute(
                """UPDATE pressure_captures
                   SET state='CANCELLED',stopped_at=?,stop_reason=?
                   WHERE id=?""",
                (utc_now(), reason, capture_id),
            )
            _event(db, capture_id, "CAPTURE_CANCELLED", reason)
            db.commit()
    finally:
        db.close()


def stop_capture(db_path: Path, capture_id: int, reason: str = "MANUAL_STOP") -> dict:
    db = connect_database(db_path)
    try:
        ensure_schema(db)
        row = db.execute("SELECT * FROM pressure_captures WHERE id=?", (capture_id,)).fetchone()
        if not row:
            raise PressureCaptureError("pressure capture not found")
        if row["state"] == "ACTIVE":
            db.execute(
                """UPDATE pressure_captures
                   SET state='STOPPED',stopped_at=?,stop_reason=?,below_since_rel=NULL
                   WHERE id=?""",
                (utc_now(), reason, capture_id),
            )
            _event(db, capture_id, "CAPTURE_STOPPED", reason)
            db.commit()
        return capture_status(db, capture_id)
    finally:
        db.close()


def _phase(relative_time_s: float, pressure_bar: float, seen_above: bool, threshold: float) -> str:
    if relative_time_s < 0:
        return "PRE_TRIGGER"
    if pressure_bar > threshold:
        return "PRESSURIZED"
    if seen_above:
        return "DECAY"
    return "POST_TRIGGER"


def ingest_edge_batch(db_path: Path, message: dict, received_at: str) -> None:
    if message.get("type") != "BATCH":
        return
    device_id = str(message.get("device_id") or "")
    if not device_id:
        return

    db = connect_database(db_path)
    try:
        ensure_schema(db)
        capture = db.execute(
            """SELECT * FROM pressure_captures
               WHERE device_id=? AND state='ACTIVE'
               ORDER BY id DESC LIMIT 1""",
            (device_id,),
        ).fetchone()
        if not capture:
            return

        capture_id = int(capture["id"])
        channels = message.get("channels") or {}
        pressure_raw = channels.get("pressure_mbar")
        if not isinstance(pressure_raw, list) or not pressure_raw:
            _event(db, capture_id, "INVALID_BATCH", "BATCH contains no pressure_mbar samples")
            db.commit()
            return
        voltage_raw = channels.get("voltage_uv") or []
        count = min(int(message["sample_count"]), len(pressure_raw))
        period_us = int(message["sample_period_us"])
        first_raw = int(message["first_sample_us"]) % MICROS_WRAP
        boot_id = str(message["boot_id"])
        sequence = int(message["sequence"])
        received_dt = _parse_stamp(received_at)
        first_wall = received_dt - timedelta(microseconds=period_us * max(0, count - 1))
        trigger_dt = _parse_stamp(capture["trigger_at"])

        wraps = int(capture["boot_wraps"] or 0)
        same_boot = capture["last_boot_id"] == boot_id
        last_ext = capture["last_extended_us"]
        last_raw = int(last_ext) % MICROS_WRAP if last_ext is not None else None

        if same_boot and last_raw is not None and first_raw < last_raw and (last_raw - first_raw) > MICROS_WRAP // 2:
            wraps += 1
            _event(db, capture_id, "ESP_TIMER_WRAP", "ESP micros() counter wrapped; timeline extended continuously")

        first_ext = first_raw + wraps * MICROS_WRAP
        if same_boot and last_ext is not None and first_ext <= int(last_ext) - MICROS_WRAP // 2:
            first_ext += MICROS_WRAP
            wraps += 1

        anchor_ext = capture["boot_anchor_extended_us"]
        anchor_rel = capture["boot_anchor_rel_s"]
        missing = 0

        if not same_boot or anchor_ext is None or anchor_rel is None:
            approx_rel = (first_wall - trigger_dt).total_seconds()
            if capture["last_boot_id"] and capture["last_boot_id"] != boot_id:
                previous_rel = capture["last_sample_rel_s"]
                if previous_rel is not None:
                    gap = max(0.0, approx_rel - float(previous_rel))
                    missing = max(0, int(round(gap * 1_000_000 / period_us)) - 1)
                _event(
                    db,
                    capture_id,
                    "ESP_REBOOT",
                    f"ESP boot id changed from {capture['last_boot_id']} to {boot_id}; capture continued",
                    missing,
                )
            wraps = 0
            first_ext = first_raw
            anchor_ext = first_ext
            anchor_rel = approx_rel
        elif same_boot and last_ext is not None:
            expected = int(last_ext) + period_us
            time_gap_samples = max(0, int(round((first_ext - expected) / period_us)))
            seq_gap_batches = 0
            if capture["last_sequence"] is not None and sequence > int(capture["last_sequence"]) + 1:
                seq_gap_batches = sequence - int(capture["last_sequence"]) - 1
            missing = max(time_gap_samples, seq_gap_batches * count)
            if missing:
                _event(
                    db,
                    capture_id,
                    "SAMPLE_GAP",
                    f"Detected telemetry gap before sequence {sequence}",
                    missing,
                )

        sample_index = int(capture["sample_count"] or 0)
        seen_above = bool(capture["seen_above"])
        below_since = capture["below_since_rel"]
        peak_bar = float(capture["peak_bar"] or 0.0)
        peak_time = capture["peak_time_s"]
        state = "ACTIVE"
        stop_reason = None
        last_rel = capture["last_sample_rel_s"]
        last_extended = capture["last_extended_us"]

        if missing and seen_above:
            below_since = None

        for index in range(count):
            raw_us = (first_raw + period_us * index) % MICROS_WRAP
            extended_us = int(first_ext) + period_us * index
            rel = float(anchor_rel) + (extended_us - int(anchor_ext)) / 1_000_000.0
            pressure_bar = float(pressure_raw[index]) / 1000.0
            voltage_v = (
                float(voltage_raw[index]) / 1_000_000.0
                if index < len(voltage_raw)
                else None
            )

            if pressure_bar > peak_bar:
                peak_bar = pressure_bar
                peak_time = rel

            if pressure_bar > float(capture["threshold_bar"]):
                seen_above = True
                below_since = None
            elif seen_above:
                if below_since is None:
                    below_since = rel
                elif (rel - float(below_since)) * 1000.0 >= int(capture["settle_ms"]):
                    state = "COMPLETED"
                    stop_reason = "PRESSURE_DECAY_BELOW_THRESHOLD"

            phase = _phase(rel, pressure_bar, seen_above, float(capture["threshold_bar"]))
            inserted = db.execute(
                """INSERT OR IGNORE INTO pressure_capture_samples(
                     capture_id,sample_index,boot_id,sequence,batch_sample_index,
                     esp_time_us,relative_time_s,pressure_bar,voltage_v,phase,received_at)
                   VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
                (
                    capture_id,
                    sample_index,
                    boot_id,
                    sequence,
                    index,
                    raw_us,
                    rel,
                    pressure_bar,
                    voltage_v,
                    phase,
                    received_at,
                ),
            )
            if inserted.rowcount:
                sample_index += 1
            last_rel = rel
            last_extended = extended_us

            if state == "COMPLETED":
                break

        db.execute(
            """UPDATE pressure_captures SET
                 state=?,
                 stopped_at=CASE WHEN ?='COMPLETED' THEN ? ELSE stopped_at END,
                 stop_reason=CASE WHEN ?='COMPLETED' THEN ? ELSE stop_reason END,
                 seen_above=?,
                 below_since_rel=?,
                 peak_bar=?,
                 peak_time_s=?,
                 sample_count=?,
                 missing_samples=missing_samples+?,
                 last_boot_id=?,
                 last_sequence=?,
                 last_extended_us=?,
                 boot_wraps=?,
                 boot_anchor_extended_us=?,
                 boot_anchor_rel_s=?,
                 last_sample_rel_s=?,
                 last_sample_received_at=?
               WHERE id=?""",
            (
                state,
                state,
                utc_now(),
                state,
                stop_reason,
                1 if seen_above else 0,
                below_since,
                peak_bar,
                peak_time,
                sample_index,
                int(missing),
                boot_id,
                sequence,
                last_extended,
                wraps,
                anchor_ext,
                anchor_rel,
                last_rel,
                received_at,
                capture_id,
            ),
        )
        if state == "COMPLETED":
            _event(
                db,
                capture_id,
                "CAPTURE_COMPLETED",
                f"Pressure remained at or below {capture['threshold_bar']:.3f} bar for {capture['settle_ms']} ms",
            )
        db.commit()
    finally:
        db.close()


def capture_status(db: sqlite3.Connection, capture_id: int) -> dict:
    ensure_schema(db)
    row = db.execute("SELECT * FROM pressure_captures WHERE id=?", (capture_id,)).fetchone()
    if not row:
        raise PressureCaptureError("pressure capture not found")
    latest = db.execute(
        """SELECT relative_time_s,pressure_bar,voltage_v,phase
           FROM pressure_capture_samples WHERE capture_id=?
           ORDER BY sample_index DESC LIMIT 1""",
        (capture_id,),
    ).fetchone()
    result = dict(row)
    result["current_pressure_bar"] = float(latest["pressure_bar"]) if latest else 0.0
    result["elapsed_s"] = max(0.0, float(latest["relative_time_s"])) if latest else 0.0
    result["latest_phase"] = latest["phase"] if latest else "WAITING"
    result["active"] = row["state"] == "ACTIVE"
    return result


def active_capture_status(db_path: Path, operation_id: str) -> dict | None:
    db = connect_database(db_path)
    try:
        ensure_schema(db)
        row = db.execute(
            """SELECT id FROM pressure_captures
               WHERE operation_id=? AND state='ACTIVE'
               ORDER BY id DESC LIMIT 1""",
            (operation_id,),
        ).fetchone()
        return capture_status(db, int(row["id"])) if row else None
    finally:
        db.close()


def list_captures(db_path: Path, operation_id: str, limit: int = 50) -> list[dict]:
    db = connect_database(db_path)
    try:
        ensure_schema(db)
        rows = db.execute(
            """SELECT id FROM pressure_captures
               WHERE operation_id=? ORDER BY id DESC LIMIT ?""",
            (operation_id, max(1, min(int(limit), 200))),
        ).fetchall()
        return [capture_status(db, int(row["id"])) for row in rows]
    finally:
        db.close()


def _duration_above(rows: Iterable[sqlite3.Row], threshold: float) -> float:
    rows = list(rows)
    total = 0.0
    for previous, current in zip(rows, rows[1:]):
        dt = float(current["relative_time_s"]) - float(previous["relative_time_s"])
        if 0 < dt <= 0.05 and float(previous["pressure_bar"]) > threshold:
            total += dt
    return total


def build_excel(db_path: Path, capture_id: int) -> tuple[io.BytesIO, str]:
    db = connect_database(db_path)
    try:
        ensure_schema(db)
        capture = db.execute(
            "SELECT * FROM pressure_captures WHERE id=?", (capture_id,)
        ).fetchone()
        if not capture:
            raise PressureCaptureError("pressure capture not found")
        samples = db.execute(
            """SELECT * FROM pressure_capture_samples
               WHERE capture_id=? ORDER BY sample_index""",
            (capture_id,),
        ).fetchall()
        events = db.execute(
            """SELECT * FROM pressure_capture_events
               WHERE capture_id=? ORDER BY id""",
            (capture_id,),
        ).fetchall()
    finally:
        db.close()

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary["A1"] = "Stellar Ops Pressure Capture"
    summary["A1"].font = Font(bold=True, size=14)

    duration_above = _duration_above(samples, float(capture["threshold_bar"]))
    summary_rows = [
        ("Capture ID", capture["id"]),
        ("State", capture["state"]),
        ("Ignition / relay time UTC", capture["trigger_at"]),
        ("Stopped UTC", capture["stopped_at"] or "ACTIVE"),
        ("Stop reason", capture["stop_reason"] or ""),
        ("Threshold (bar)", float(capture["threshold_bar"])),
        ("Peak pressure (bar)", float(capture["peak_bar"] or 0.0)),
        ("Peak time from relay (s)", capture["peak_time_s"] if capture["peak_time_s"] is not None else ""),
        ("Time above 0.5 bar (s)", round(duration_above, 6)),
        ("Recorded samples", int(capture["sample_count"] or 0)),
        ("Detected missing samples", int(capture["missing_samples"] or 0)),
        ("Nominal sample rate (Hz)", float(capture["sample_rate_hz"])),
        ("ESP device", capture["device_id"]),
    ]
    for row_index, (label, value) in enumerate(summary_rows, start=3):
        summary.cell(row=row_index, column=1, value=label).font = Font(bold=True)
        summary.cell(row=row_index, column=2, value=value)

    event_start = 18
    summary.cell(row=event_start, column=1, value="Detected events").font = Font(bold=True)
    for offset, event_row in enumerate(events, start=1):
        summary.cell(row=event_start + offset, column=1, value=event_row["occurred_at"])
        summary.cell(row=event_start + offset, column=2, value=event_row["event_type"])
        summary.cell(row=event_start + offset, column=3, value=event_row["detail"])
        summary.cell(row=event_start + offset, column=4, value=int(event_row["missing_samples"] or 0))

    data = wb.create_sheet("Pressure data")
    headers = [
        "Sample",
        "Time from relay (s)",
        "Pressure (bar)",
        "Voltage (V)",
        "Phase",
        "Boot ID",
        "Sequence",
        "ESP time (us)",
        "Received UTC",
    ]
    for column, header in enumerate(headers, start=1):
        cell = data.cell(row=1, column=column, value=header)
        cell.font = Font(bold=True)

    for excel_row, sample in enumerate(samples, start=2):
        values = [
            int(sample["sample_index"]),
            float(sample["relative_time_s"]),
            float(sample["pressure_bar"]),
            float(sample["voltage_v"]) if sample["voltage_v"] is not None else None,
            sample["phase"],
            sample["boot_id"],
            int(sample["sequence"]),
            int(sample["esp_time_us"]),
            sample["received_at"],
        ]
        for column, value in enumerate(values, start=1):
            data.cell(row=excel_row, column=column, value=value)

    data.freeze_panes = "A2"
    widths = [12, 20, 18, 16, 18, 24, 12, 18, 28]
    for index, width in enumerate(widths, start=1):
        data.column_dimensions[get_column_letter(index)].width = width
    summary.column_dimensions["A"].width = 30
    summary.column_dimensions["B"].width = 28
    summary.column_dimensions["C"].width = 70

    if samples:
        chart = LineChart()
        chart.title = "Pressure vs Time"
        chart.y_axis.title = "Pressure (bar)"
        chart.x_axis.title = "Time from relay (s)"
        pressure_ref = Reference(data, min_col=3, min_row=1, max_row=len(samples) + 1)
        time_ref = Reference(data, min_col=2, min_row=2, max_row=len(samples) + 1)
        chart.add_data(pressure_ref, titles_from_data=True)
        chart.set_categories(time_ref)
        chart.height = 12
        chart.width = 24
        summary.add_chart(chart, "F3")

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f"pressure-capture-{capture_id}.xlsx"
    return output, filename
