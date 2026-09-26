import io
import json
import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from stellar_ops.edge_gateway import database
from stellar_ops import control as control_module
from stellar_ops.pressure_capture import (
    PressureCaptureError,
    build_excel,
    capture_status,
    ensure_schema,
    ingest_edge_batch,
    start_capture,
    stop_capture,
)


def stamp():
    return datetime.now(timezone.utc).isoformat(timespec="milliseconds")


def batch_message(sequence, first_sample_us, pressures_mbar, boot_id="boot-1"):
    return {
        "protocol": "SMTCS-EDGE/1",
        "type": "BATCH",
        "device_id": "PT-01",
        "boot_id": boot_id,
        "sequence": sequence,
        "first_sample_us": first_sample_us,
        "sample_period_us": 5000,
        "sample_count": len(pressures_mbar),
        "channels": {
            "pressure_mbar": pressures_mbar,
            "voltage_uv": [1000000 + i for i in range(len(pressures_mbar))],
        },
    }


class PressureCaptureTests(unittest.TestCase):
    def make_db(self):
        directory = tempfile.TemporaryDirectory()
        path = Path(directory.name) / "control.db"
        with database(path) as db:
            ensure_schema(db)
            now = stamp()
            db.execute(
                """INSERT INTO edge_sessions(
                     device_id,boot_id,remote_addr,firmware,connected_at,last_seen,status)
                   VALUES('PT-01','boot-1','192.168.1.50:50000','test',?,?, 'STREAMING')""",
                (now, now),
            )
            db.execute(
                """INSERT INTO edge_batches(
                     device_id,boot_id,sequence,received_at,first_sample_us,
                     sample_period_us,sample_count,channels_json)
                   VALUES('PT-01','boot-1',0,?,1000000,5000,20,?)""",
                (
                    now,
                    json.dumps(
                        {
                            "pressure_mbar": [0] * 20,
                            "voltage_uv": [1000000] * 20,
                        }
                    ),
                ),
            )
            db.commit()
        return directory, path

    def test_capture_starts_with_pretrigger_and_auto_stops_after_decay(self):
        directory, path = self.make_db()
        try:
            started = start_capture(path, "OP-001")
            self.assertEqual(started["state"], "ACTIVE")
            self.assertGreater(started["sample_count"], 0)

            ingest_edge_batch(path, batch_message(1, 1100000, [1200] * 20), stamp())
            ingest_edge_batch(path, batch_message(2, 1200000, [200] * 20), stamp())
            ingest_edge_batch(path, batch_message(3, 1300000, [200] * 20), stamp())
            ingest_edge_batch(path, batch_message(4, 1400000, [200] * 20), stamp())

            with database(path) as db:
                final = capture_status(db, started["id"])
            self.assertEqual(final["state"], "COMPLETED")
            self.assertEqual(final["stop_reason"], "PRESSURE_DECAY_BELOW_THRESHOLD")
            self.assertGreaterEqual(final["peak_bar"], 1.2)
            self.assertTrue(final["seen_above"])
        finally:
            directory.cleanup()

    def test_no_pressure_rise_requires_manual_stop(self):
        directory, path = self.make_db()
        try:
            started = start_capture(path, "OP-001")
            for sequence in range(1, 7):
                ingest_edge_batch(
                    path,
                    batch_message(sequence, 1000000 + sequence * 100000, [100] * 20),
                    stamp(),
                )
            with database(path) as db:
                still_active = capture_status(db, started["id"])
            self.assertEqual(still_active["state"], "ACTIVE")
            stopped = stop_capture(path, started["id"])
            self.assertEqual(stopped["state"], "STOPPED")
            self.assertEqual(stopped["stop_reason"], "MANUAL_STOP")
        finally:
            directory.cleanup()

    def test_detects_missing_samples_from_esp_clock_even_without_sequence_gap(self):
        directory, path = self.make_db()
        try:
            started = start_capture(path, "OP-001")
            ingest_edge_batch(path, batch_message(1, 1100000, [800] * 20), stamp())
            # Sequence is continuous but the ESP sample clock jumps by 200 ms
            # instead of 100 ms. This models firmware dropping one unsent batch.
            ingest_edge_batch(path, batch_message(2, 1300000, [800] * 20), stamp())
            with database(path) as db:
                status = capture_status(db, started["id"])
                gap = db.execute(
                    """SELECT * FROM pressure_capture_events
                       WHERE capture_id=? AND event_type='SAMPLE_GAP'
                       ORDER BY id DESC LIMIT 1""",
                    (started["id"],),
                ).fetchone()
            self.assertIsNotNone(gap)
            self.assertGreaterEqual(status["missing_samples"], 20)
        finally:
            directory.cleanup()

    def test_capture_survives_boot_id_change_and_exports_excel(self):
        directory, path = self.make_db()
        try:
            started = start_capture(path, "OP-001")
            ingest_edge_batch(path, batch_message(1, 1100000, [1000] * 20), stamp())
            ingest_edge_batch(
                path,
                batch_message(0, 10000, [100] * 20, boot_id="boot-2"),
                stamp(),
            )
            stop_capture(path, started["id"])
            output, filename = build_excel(path, started["id"])
            self.assertTrue(filename.endswith(".xlsx"))
            workbook = load_workbook(io.BytesIO(output.getvalue()))
            self.assertIn("Summary", workbook.sheetnames)
            self.assertIn("Pressure data", workbook.sheetnames)
            self.assertGreater(workbook["Pressure data"].max_row, 1)
            with database(path) as db:
                reboot = db.execute(
                    """SELECT * FROM pressure_capture_events
                       WHERE capture_id=? AND event_type='ESP_REBOOT'""",
                    (started["id"],),
                ).fetchone()
            self.assertIsNotNone(reboot)
        finally:
            directory.cleanup()

    def test_dynamic_bench_relay_arms_capture_before_sending_on(self):
        directory, path = self.make_db()
        original_db = control_module.CONTROL_DB
        try:
            control_module.CONTROL_DB = path
            with patch(
                "stellar_ops.edge_runtime.send_bench_led_state",
                return_value={"ok": True, "device_id": "PT-01", "state": "ON"},
            ) as send:
                result = control_module._execute_bench_led_set("PT-01", True)
            self.assertTrue(result["ok"])
            self.assertIn("capture", result)
            send.assert_called_once_with(device_id="PT-01", on=True)
            with database(path) as db:
                active = db.execute(
                    "SELECT * FROM pressure_captures WHERE state='ACTIVE'"
                ).fetchone()
            self.assertIsNotNone(active)
        finally:
            control_module.CONTROL_DB = original_db
            directory.cleanup()

    def test_stale_or_missing_telemetry_blocks_capture(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "control.db"
            with database(path) as db:
                ensure_schema(db)
            with self.assertRaises(PressureCaptureError):
                start_capture(path, "OP-001")


if __name__ == "__main__":
    unittest.main()
